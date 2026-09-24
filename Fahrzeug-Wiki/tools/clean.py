"""Erzeugt die bereinigte Referenztabelle aus der unveränderten Rohquelle + tools/corrections.json.

Aufruf:  python tools/clean.py
Ausgabe: data/tn_batterycheck_bereinigt.xlsx  (Blätter: Fahrzeuge, Änderungsprotokoll, Offene Klärungen, Legende)
         data/tn_batterycheck_bereinigt.csv   (UTF-8, Komma-getrennt, Dezimalpunkt)
Die Rohquelle raw/tn_batterycheck_alle_daten.xlsx wird nur gelesen, nie verändert.
"""
import csv
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw" / "tn_batterycheck_alle_daten.xlsx"
CORR = ROOT / "tools" / "corrections.json"
OUT = ROOT / "data"
REPO = "https://github.com/Thitronik01/Frank-Repo/issues/"

COLUMNS = ["ID", "Zeile (Rohquelle)", "Marke", "Modell", "Brutto kWh", "Netto kWh", "Puffer kWh", "Puffer %",
           "Status", "Alias von Zeile", "Modelljahr von", "Modelljahr bis", "Batterie-Generation", "Bemerkung",
           "Quelle", "Issue", "Marke (Rohquelle)", "Modell (Rohquelle)", "Brutto (Rohquelle)", "Netto (Rohquelle)"]
STATUS_FILL = {"korrigiert": "E1F0EA", "bestätigt": "F1F8F4", "zu prüfen": "FCEFD2", "Näherungswert": "FCEFD2", "Alias": "EFEDE7"}


def kwh(value):
    s = str(value).strip()
    return float(re.search(r"\d+(?:\.\d+)?", s).group()), s.lower().startswith("ca.")


def model_name(model, brand_raw, c):
    name = c["modelle"].get(model, model)
    for rule in c["modell_regeln"]:
        if rule["marke"] == brand_raw:
            name = re.sub(rule["muster"], rule["ersatz"], name)
    return name


def main():
    c = json.loads(CORR.read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(RAW, data_only=True).worksheets[0]
    remove = {e["zeile"]: e for e in c["entfernen"]}
    values = {e["zeile"]: e for e in c["werte"]}
    names = {e["zeile"]: e for e in c.get("namen_zeile", [])}
    alias = {e["zeile"]: e for e in c["aliase"]}
    flags = {}
    for e in c["pruefen"]:
        for z in e["zeilen"]:
            flags.setdefault(z, []).append(e)

    rows, log = [], []
    for zeile, (marke, modell, brutto_raw, netto_raw) in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        brutto, ca_b = kwh(brutto_raw)
        netto, ca_n = kwh(netto_raw)
        if zeile in remove:
            e = remove[zeile]
            log.append([c["stand"], zeile, f"{marke} {modell}", "Zeile", f"{brutto_raw} / {netto_raw}", "entfernt",
                        e["grund"], "", e.get("issue", "")])
            continue
        rec = {"ID": f"Z{zeile:03d}", "Zeile (Rohquelle)": zeile, "Marke": c["marken"].get(marke, marke),
               "Modell": model_name(modell, marke, c), "Brutto kWh": brutto, "Netto kWh": netto, "Status": "unverändert",
               "Alias von Zeile": "", "Modelljahr von": "", "Modelljahr bis": "", "Batterie-Generation": "",
               "Bemerkung": "", "Quelle": "Rohquelle", "Issue": "",
               "Marke (Rohquelle)": marke, "Modell (Rohquelle)": modell, "Brutto (Rohquelle)": brutto_raw,
               "Netto (Rohquelle)": netto_raw}
        notes, issues = [], []
        if ca_b or ca_n:
            rec["Status"] = "Näherungswert"
            notes.append("In der Rohquelle als „ca.“ angegeben.")
        if zeile in names:
            e = names[zeile]
            log.append([c["stand"], zeile, f"{marke} {modell}", "Modell", modell, e["modell"], e["grund"],
                        " ".join(e.get("quelle", [])), e.get("issue", "")])
            rec["Modell"] = e["modell"]
            if e.get("issue"):
                issues.append(e["issue"])
        if zeile in values:
            e = values[zeile]
            for field, key in (("Brutto kWh", "brutto"), ("Netto kWh", "netto")):
                if key in e and e[key] != rec[field]:
                    log.append([c["stand"], zeile, f"{marke} {modell}", field, rec[field], e[key], e["grund"],
                                " ".join(e.get("quelle", [])), e.get("issue", "")])
                    rec[field] = e[key]
            for k in ("Modelljahr von", "Modelljahr bis", "Batterie-Generation"):
                if e.get(k):
                    rec[k] = e[k]
            rec["Status"] = e.get("status", "korrigiert")
            rec["Quelle"] = " ".join(e.get("quelle", [])) or rec["Quelle"]
            notes.append(e["grund"])
            if e.get("issue"):
                issues.append(e["issue"])
        if zeile in alias:
            e = alias[zeile]
            rec["Status"] = "Alias" if rec["Status"] == "unverändert" else rec["Status"]
            rec["Alias von Zeile"] = e["von"]
            notes.append(e["grund"])
            issues.append(e.get("issue", ""))
        for e in flags.get(zeile, []):
            if rec["Status"] in ("unverändert", "Alias"):
                rec["Status"] = "zu prüfen"
            notes.append(e["bemerkung"])
            issues.append(e["issue"])
        rec["Bemerkung"] = " ".join(dict.fromkeys(n for n in notes if n))
        rec["Issue"] = " ".join(f"#{i}" for i in dict.fromkeys(i for i in issues if i))
        rec["Puffer kWh"] = round(rec["Brutto kWh"] - rec["Netto kWh"], 2)
        rec["Puffer %"] = round(rec["Puffer kWh"] / rec["Brutto kWh"] * 100, 1)
        if rec["Netto kWh"] > rec["Brutto kWh"]:
            raise SystemExit(f"Zeile {zeile}: Netto > Brutto nach Bereinigung")
        rows.append(rec)

    # Namensänderungen protokollieren (nur echte Änderungen)
    for r in rows:
        if r["Modell"] != r["Modell (Rohquelle)"] and r["Zeile (Rohquelle)"] not in names:
            log.append([c["stand"], r["Zeile (Rohquelle)"], f'{r["Marke (Rohquelle)"]} {r["Modell (Rohquelle)"]}', "Modell",
                        r["Modell (Rohquelle)"], r["Modell"], "Schreibweise vereinheitlicht (Herstellerschreibweise)", "", ""])
    for r in {(r["Marke (Rohquelle)"], r["Marke"]) for r in rows if r["Marke"] != r["Marke (Rohquelle)"]}:
        log.append([c["stand"], "alle", r[0], "Marke", r[0], r[1], "Herstellerschreibweise", "", ""])

    OUT.mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = "Fahrzeuge"
    sh.append(COLUMNS)
    for r in rows:
        sh.append([r[k] if k not in ("Puffer kWh", "Puffer %") else r[k] for k in COLUMNS])
    style(sh, {"Brutto kWh": "0.0#", "Netto kWh": "0.0#", "Puffer kWh": "0.0#", "Puffer %": "0.0"})
    status_col = COLUMNS.index("Status") + 1
    for row in sh.iter_rows(min_row=2):
        fill = STATUS_FILL.get(row[status_col - 1].value)
        if fill:
            for cell in row[:COLUMNS.index("Issue") + 1]:
                cell.fill = PatternFill("solid", fgColor=fill)
        issue_cell = row[COLUMNS.index("Issue")]
        m = re.match(r"#(\d+)", str(issue_cell.value or ""))
        if m:
            issue_cell.hyperlink = REPO + m.group(1)
            issue_cell.font = Font(color="0A6B4E", underline="single")

    lg = wb.create_sheet("Änderungsprotokoll")
    lg.append(["Datum", "Zeile (Rohquelle)", "Fahrzeug (Rohquelle)", "Feld", "Alt", "Neu", "Grund", "Quelle", "Issue"])
    for entry in log:
        lg.append([*entry[:8], f"#{entry[8]}" if entry[8] else ""])
    style(lg, {})

    op = wb.create_sheet("Offene Klärungen")
    op.append(["Issue", "Thema", "Betroffene Zeilen", "Stand"])
    for e in c["offene_issues"]:
        op.append([f"#{e['issue']}", e["thema"], e["zeilen"], e["stand"]])
    style(op, {})
    for row in op.iter_rows(min_row=2):
        row[0].hyperlink = REPO + row[0].value.lstrip("#")
        row[0].font = Font(color="0A6B4E", underline="single")

    le = wb.create_sheet("Legende")
    for line in c["legende"]:
        le.append(line)
    style(le, {})
    wb.save(OUT / "tn_batterycheck_bereinigt.xlsx")

    with open(OUT / "tn_batterycheck_bereinigt.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow([r[k] for k in COLUMNS])

    counts = {}
    for r in rows:
        counts[r["Status"]] = counts.get(r["Status"], 0) + 1
    print(f"{len(rows)} Zeilen ({469 - len(rows)} entfernt) · {counts} · {len(log)} Protokolleinträge")


def style(sheet, numfmt):
    head = [c.value for c in sheet[1]]
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A6B4E")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, h in enumerate(head, start=1):
        width = max([len(str(h or ""))] + [len(str(c.value or "")) for c in sheet[get_column_letter(i)][1:200]])
        sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        if h in numfmt:
            for c in sheet[get_column_letter(i)][1:]:
                c.number_format = numfmt[h]


if __name__ == "__main__":
    main()
